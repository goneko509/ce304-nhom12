"""Kiem chung hau kiem nhan heuristic tren MAU su kien (stratified random sampling).

Buoc 1 - rut mau + ve do thi + xuat file duyet:
    python kiem_chung_mau.py sample
    -> thongke_output/kiem_chung/kiem_chung_mau.xlsx   (nguoi lai dien cot ket_qua)
    -> thongke_output/kiem_chung/plots/S###.png         (tin hieu quanh moi su kien)

Buoc 2 - sau khi dien xong ket_qua (DUNG / SAI / KHONG_CHAC):
    python kiem_chung_mau.py evaluate
    -> thongke_output/kiem_chung/ket_qua_kiem_chung.txt
    -> thongke_output/kiem_chung/bang_kiem_chung.tex    (bang LaTeX dua vao Chuong 4)

Dau vao: <trip>/<trip>_main_events.csv do thong_ke_chuyen_di.py sinh ra, va
<trip>/detected_events_unified/merged_motion_features.csv (tin hieu da tien xu ly).
Mau co seed co dinh nen chay lai 'sample' cho dung cung bo su kien.
"""

import argparse
import math
import os
import re
from datetime import date

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from thong_ke_chuyen_di import RAW_ROOT, trip_start_time  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "thongke_output", "kiem_chung")
PLOTS = os.path.join(OUT, "plots")
XLSX = os.path.join(OUT, "kiem_chung_mau.xlsx")
SEED = 42
CONTEXT_SEC = 5.0  # hien thi tin hieu truoc/sau su kien
N_PER_STRATUM = 5  # 6 tang x 5 = 30 mau (giam tu 100 de kip duyet ky tung mau)

# Tang mau: (ten tang, dieu kien, so mau muc tieu, cau hoi kiem chung)
STRATA = [
    ("BRAKE_HARD", lambda d: (d.event_type == "BRAKE") & (d.severity == "HARD"), N_PER_STRATUM,
     "Có phải PHANH GẤP thật (giảm tốc mạnh, đột ngột)?"),
    ("BRAKE_MODERATE", lambda d: (d.event_type == "BRAKE") & (d.severity == "MODERATE"), N_PER_STRATUM,
     "Có phải GIẢM TỐC rõ rệt (mức vừa)?"),
    ("ACCEL_HARD", lambda d: (d.event_type == "ACCEL") & (d.severity == "HARD"), N_PER_STRATUM,
     "Có phải TĂNG TỐC GẤP thật?"),
    ("ACCEL_MODERATE", lambda d: (d.event_type == "ACCEL") & (d.severity == "MODERATE"), N_PER_STRATUM,
     "Có phải TĂNG TỐC rõ rệt (mức vừa)?"),
    ("STEER", lambda d: d.event_type == "STEER", N_PER_STRATUM,
     "Có phải ĐÁNH LÁI / VÀO CUA thật?"),
    ("COMBO", lambda d: d.event_type.str.contains(r"\+"), N_PER_STRATUM,
     "Có đúng là phanh/tăng tốc KÈM đánh lái?"),
]
RESULT_CHOICES = ["DUNG", "SAI", "KHONG_CHAC"]
TRUE_LABEL_CHOICES = ["BRAKE", "ACCEL", "STEER", "NORMAL", "STOP", "NHIEU_CAM_BIEN", "KHAC"]

INK, INK2, SURFACE = "#0b0b0b", "#52514e", "#fcfcfb"
SERIES, SHADE = "#2a78d6", "#cde2fb"

SAMPLE_COLUMNS = ["sample_id", "tang_mau", "trip", "timestamp", "event_type", "severity",
                  "heuristic_label", "speed_group", "start_time", "end_time", "duration_s",
                  "speed_start_kmh", "speed_end_kmh", "min_accel", "max_accel", "max_gyr_z",
                  "cau_hoi", "do_thi", "ban_do_trip", "google_maps",
                  "ket_qua", "nhan_dung_neu_sai", "ghi_chu", "nguoi_kiem_chung", "ngay_kiem_chung"]


def load_all_main_events(root):
    frames = []
    for trip in sorted(os.listdir(root)):
        if not re.fullmatch(r"\d{8}_\d{6}", trip):
            continue
        path = os.path.join(root, trip, f"{trip}_main_events.csv")
        if not os.path.isfile(path):
            raise FileNotFoundError(f"Chua co {path} - hay chay thong_ke_chuyen_di.py truoc.")
        df = pd.read_csv(path, encoding="utf-8-sig")
        df.insert(0, "trip", trip)
        frames.append(df)
    return pd.concat(frames, ignore_index=True)


def draw_sample_plots(root, samples):
    os.makedirs(PLOTS, exist_ok=True)
    cols = ["time", "speed_kmh", "acc_y_smooth", "jerk", "gyr_z"]
    panels = [("speed_kmh", "Tốc độ (km/h)", None),
              ("acc_y_smooth", "Gia tốc dọc a_y (m/s²)", (-1.2, 1.2)),
              ("jerk", "Độ giật j (m/s³)", (-3.0, 3.0)),
              ("gyr_z", "Vận tốc góc ω_z (rad/s)", None)]
    for trip, group in samples.groupby("trip"):
        feat = pd.read_csv(os.path.join(root, trip, "detected_events_unified", "merged_motion_features.csv"),
                           usecols=cols, encoding="utf-8-sig")
        t = feat["time"].to_numpy()
        for _, ev in group.iterrows():
            a, b = ev.start_time - CONTEXT_SEC, ev.end_time + CONTEXT_SEC
            seg = feat[(t >= a) & (t <= b)]
            fig, axes = plt.subplots(4, 1, figsize=(8, 7.2), sharex=True, facecolor=SURFACE)
            for ax, (col, ylabel, thr) in zip(axes, panels):
                ax.set_facecolor(SURFACE)
                ax.axvspan(ev.start_time, ev.end_time, color=SHADE, alpha=0.7, lw=0)
                if thr:
                    for y in thr:
                        ax.axhline(y, color=INK2, lw=0.8, ls=(0, (4, 3)))
                ax.axhline(0, color=INK2, lw=0.6, alpha=0.5)
                ax.plot(seg["time"], seg[col], color=SERIES, lw=1.4)
                if col == "jerk" and len(seg):
                    # gai jerk do nhieu cam bien (+-hang tram m/s3) che mat nguong +-3:
                    # gioi han truc theo phan vi 95 cua |j|, ghi chu ro tren nhan truc
                    lim = max(10.0, float(np.nanpercentile(np.abs(seg[col]), 95)) * 1.5)
                    ax.set_ylim(-lim, lim)
                    ylabel += "\n(cắt trục)"
                ax.set_ylabel(ylabel, color=INK2, fontsize=8.5)
                ax.tick_params(colors=INK2, labelsize=8)
                ax.grid(True, axis="y", color="#e6e5e1", lw=0.6)
                for side in ("top", "right"):
                    ax.spines[side].set_visible(False)
                for side in ("left", "bottom"):
                    ax.spines[side].set_color("#c9c8c3")
            axes[-1].set_xlabel("Thời gian trong chuyến (s) — vùng tô: cửa sổ sự kiện heuristic; nét đứt: ngưỡng luật",
                                color=INK2, fontsize=8.5)
            fig.suptitle(f"{ev.sample_id} · {ev.tang_mau} · {ev.heuristic_label} · trip {trip} · {ev.timestamp}",
                         color=INK, fontsize=10, x=0.01, ha="left")
            fig.tight_layout(rect=(0, 0, 1, 0.97))
            fig.savefig(os.path.join(PLOTS, f"{ev.sample_id}.png"), dpi=110, facecolor=SURFACE)
            plt.close(fig)


def write_review_xlsx(samples):
    wb = Workbook()
    ws = wb.active
    ws.title = "kiem_chung"
    ws.append(SAMPLE_COLUMNS)
    for _, r in samples.iterrows():
        ws.append([r.get(c, "") for c in SAMPLE_COLUMNS])
    head_fill = PatternFill("solid", fgColor="E8E7E3")
    edit_fill = PatternFill("solid", fgColor="FFF6D6")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = head_fill
    idx = {c: i + 1 for i, c in enumerate(SAMPLE_COLUMNS)}
    n = len(samples) + 1
    for row in range(2, n + 1):
        for c in ("do_thi", "ban_do_trip", "google_maps"):
            cell = ws.cell(row=row, column=idx[c])
            if cell.value:
                cell.hyperlink = cell.value
                cell.value = {"do_thi": "xem đồ thị", "ban_do_trip": "bản đồ trip", "google_maps": "Google Maps"}[c]
                cell.font = Font(color="1C5CAB", underline="single")
        for c in ("ket_qua", "nhan_dung_neu_sai", "ghi_chu", "nguoi_kiem_chung", "ngay_kiem_chung"):
            ws.cell(row=row, column=idx[c]).fill = edit_fill
    letter = lambda c: ws.cell(row=1, column=idx[c]).column_letter  # noqa: E731
    dv1 = DataValidation(type="list", formula1='"' + ",".join(RESULT_CHOICES) + '"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1='"' + ",".join(TRUE_LABEL_CHOICES) + '"', allow_blank=True)
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    dv1.add(f"{letter('ket_qua')}2:{letter('ket_qua')}{n}")
    dv2.add(f"{letter('nhan_dung_neu_sai')}2:{letter('nhan_dung_neu_sai')}{n}")
    widths = {"cau_hoi": 42, "timestamp": 24, "heuristic_label": 22, "ghi_chu": 40, "trip": 17}
    for c, i in idx.items():
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = widths.get(c, 13)
    ws.freeze_panes = "D2"
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column == idx["cau_hoi"])

    guide = wb.create_sheet("huong_dan")
    for line in [
        "HƯỚNG DẪN KIỂM CHỨNG HẬU KIỂM (người lái điền)",
        "1. Mở 'xem đồ thị': vùng tô xanh là cửa sổ sự kiện do luật heuristic phát hiện; nét đứt là ngưỡng luật.",
        "2. Đối chiếu với kịch bản chạy, bản đồ trip, Google Maps (vị trí) và trí nhớ của người lái.",
        "3. Cột ket_qua: DUNG = đúng loại sự kiện; SAI = không phải sự kiện đó; KHONG_CHAC = không đủ căn cứ.",
        "   Quy ước: sự kiện thật xảy ra trong khoảng ±2 s quanh vùng tô vẫn tính là DUNG (GPS chỉ cập nhật ~1 Hz nên"
        " cửa sổ heuristic có thể lệch thời gian); ghi 'lech thoi gian' vào ghi_chu.",
        "4. Nếu SAI, chọn nhan_dung_neu_sai (BRAKE/ACCEL/STEER/NORMAL/STOP/NHIEU_CAM_BIEN/KHAC).",
        "5. Ghi tên người kiểm chứng và ngày kiểm chứng (kiểm chứng HẬU KIỂM, không phải nhãn ghi lúc lái).",
        f"Mẫu ngẫu nhiên phân tầng, seed = {SEED}; chạy lại 'python kiem_chung_mau.py sample' cho cùng bộ mẫu.",
    ]:
        guide.append([line])
    guide.column_dimensions["A"].width = 120
    wb.save(XLSX)


def cmd_sample(root):
    events = load_all_main_events(root)
    rng = np.random.default_rng(SEED)
    parts = []
    for name, cond, k, question in STRATA:
        pool = events[cond(events)]
        take = pool if len(pool) <= k else pool.iloc[np.sort(rng.choice(len(pool), size=k, replace=False))]
        take = take.assign(tang_mau=name, cau_hoi=question)
        parts.append(take)
        print(f"  {name:<15} quan the {len(pool):>4}  ->  mau {len(take):>3}")
    samples = pd.concat(parts).sort_values(["trip", "start_time"]).reset_index(drop=True)
    samples.insert(0, "sample_id", [f"S{i:03d}" for i in range(1, len(samples) + 1)])
    samples["do_thi"] = samples["sample_id"].map(lambda s: f"plots/{s}.png")
    samples["ban_do_trip"] = samples["trip"].map(
        lambda tr: "file:///" + os.path.join(root, tr, "detected_events_unified", "driving_events_map.html").replace("\\", "/"))
    samples["google_maps"] = [f"https://www.google.com/maps?q={la},{lo}" for la, lo in zip(samples.lat, samples.lon)]
    for c in ("ket_qua", "nhan_dung_neu_sai", "ghi_chu", "nguoi_kiem_chung", "ngay_kiem_chung"):
        samples[c] = ""
    os.makedirs(OUT, exist_ok=True)
    if os.path.isfile(XLSX):
        filled = pd.read_excel(XLSX)
        if filled["ket_qua"].notna().any():
            raise SystemExit(f"{XLSX} da co ket qua kiem chung - khong ghi de. Doi ten file cu neu muon tao lai.")
    print("Dang ve do thi...")
    draw_sample_plots(root, samples)
    write_review_xlsx(samples)
    print(f"Da tao {len(samples)} mau: {XLSX}")


def wilson(k, n, z=1.96):
    if n == 0:
        return float("nan"), float("nan")
    p = k / n
    den = 1 + z * z / n
    centre = (p + z * z / (2 * n)) / den
    half = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / den
    return centre - half, centre + half


def cmd_evaluate():
    df = pd.read_excel(XLSX)
    df["ket_qua"] = df["ket_qua"].fillna("").astype(str).str.strip().str.upper()
    pending = (df["ket_qua"] == "").sum()
    rows, lines = [], []
    for name, *_ in STRATA + [("TONG", None, None, None)]:
        d = df if name == "TONG" else df[df["tang_mau"] == name]
        n_d, n_s, n_u = (d["ket_qua"] == "DUNG").sum(), (d["ket_qua"] == "SAI").sum(), (d["ket_qua"] == "KHONG_CHAC").sum()
        decided = n_d + n_s
        p = n_d / decided if decided else float("nan")
        lo, hi = wilson(n_d, decided)
        rows.append((name, len(d), n_d, n_s, n_u, p, lo, hi))
    lines.append(f"KET QUA KIEM CHUNG HAU KIEM NHAN HEURISTIC - {date.today():%d/%m/%Y}")
    lines.append(f"Tong mau: {len(df)} | chua dien: {pending}")
    reviewers = ", ".join(sorted(set(df["nguoi_kiem_chung"].dropna().astype(str))))
    lines.append(f"Nguoi kiem chung: {reviewers or '(chua ghi)'}")
    lines.append("")
    lines.append(f"{'Tang':<15}{'Mau':>5}{'Dung':>6}{'Sai':>5}{'K.chac':>8}{'Do chinh xac':>14}{'KTC 95% (Wilson)':>20}")
    for name, n, d, s, u, p, lo, hi in rows:
        ps = "-" if math.isnan(p) else f"{100 * p:.1f}%"
        ci = "-" if math.isnan(lo) else f"[{100 * lo:.1f}; {100 * hi:.1f}]%"
        lines.append(f"{name:<15}{n:>5}{d:>6}{s:>5}{u:>8}{ps:>14}{ci:>20}")
    wrong = df[df["ket_qua"] == "SAI"]
    if len(wrong):
        lines.append("")
        lines.append("Nhan dung cua cac su kien SAI (ma tran nham lan rut gon):")
        ct = pd.crosstab(wrong["tang_mau"], wrong["nhan_dung_neu_sai"].fillna("(trong)"))
        lines.extend("   " + x for x in ct.to_string().split("\n"))
    text = "\n".join(lines)
    with open(os.path.join(OUT, "ket_qua_kiem_chung.txt"), "w", encoding="utf-8") as f:
        f.write(text + "\n")
    print(text)

    def pct(x):
        return "---" if math.isnan(x) else f"{100 * x:.1f}".replace(".", "{,}") + r"\%"
    tex = [r"\begin{table}[H]", r"\centering", r"\setlength{\tabcolsep}{4pt}",
           rf"\caption{{Kiểm chứng hậu kiểm nhãn heuristic trên mẫu {len(df)} sự kiện (ngẫu nhiên phân tầng, seed {SEED}).}}",
           r"\vspace{0.2cm}", r"\label{tab:label_verification}",
           r"\begin{tabular}{|l|c|c|c|c|c|c|}", r"\hline",
           r"\textbf{Tầng mẫu} & \textbf{Mẫu} & \textbf{Đúng} & \textbf{Sai} & \textbf{Không chắc} & \textbf{Độ chính xác} & \textbf{KTC 95\%} \\ \hline"]
    for name, n, d, s, u, p, lo, hi in rows:
        label = r"\textbf{Tổng}" if name == "TONG" else r"\texttt{" + name.replace("_", r"\_") + "}"
        ci = "---" if math.isnan(lo) else f"{pct(lo)}--{pct(hi)}"
        tex.append(f"{label} & {n} & {d} & {s} & {u} & {pct(p)} & {ci} \\\\ \\hline")
    tex += [r"\end{tabular}", r"\end{table}"]
    with open(os.path.join(OUT, "bang_kiem_chung.tex"), "w", encoding="utf-8") as f:
        f.write("\n".join(tex) + "\n")
    print(f"\nDa ghi: {os.path.join(OUT, 'ket_qua_kiem_chung.txt')}\n        {os.path.join(OUT, 'bang_kiem_chung.tex')}")


def main():
    ap = argparse.ArgumentParser(description="Kiem chung hau kiem nhan heuristic tren mau su kien")
    ap.add_argument("command", choices=["sample", "evaluate"])
    ap.add_argument("--root", default=RAW_ROOT)
    args = ap.parse_args()
    if args.command == "sample":
        cmd_sample(args.root)
    else:
        cmd_evaluate()


if __name__ == "__main__":
    main()
